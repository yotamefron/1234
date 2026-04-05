#!/usr/bin/env python3
"""Build ametrine_inventory.xlsx per specification."""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

OUT_PATH = "/home/user/1234/ametrine_inventory.xlsx"

# ── DATA ──────────────────────────────────────────────────────────────
PRINTS = [
    ("Desert",              "מדברי",            ""),
    ("Desert-Green",        "מדברי",            ""),
    ("Desert Coyote",       "מדברי",            ""),
    ("Sand",                "מדברי",            ""),
    ("Ever-Green",          "מיוער",            ""),
    ("Forest Green",        "מיוער",            ""),
    ("Kestrel",             "מיוער",            ""),
    ("Woodland",            "מיוער",            ""),
    ("Woodland Brown",      "מיוער",            ""),
    ("Urban Green",         "מיוער",            "בולדר/שטח בנוי"),
    ("Snirt Snow",          "שלג",              ""),
    ("Fresh Snow",          "שלג",              ""),
    ("Marine",              "ימי",              ""),
    ("Urban",               "בולדר/שטח בנוי",   ""),
    ("Boulder",             "בולדר/שטח בנוי",   ""),
    ("NATO",                "מיוער",            ""),
    ("Kestrel (אמריקאי)",   "מיוער",            ""),
]  # 17 → rows 2-18

ENVIRONMENTS = ["הכל","מדברי","מיוער","שלג","ימי","בולדר/שטח בנוי","אחר"]  # 7

PRODUCT_LINES = ["הכל","OverGarment","Hide Site","Platform Hide Site",
                  "Uniform","Blankets","Urban","Platform On-The-Move","Accessories"]  # 9

PRODUCTS = [
    ("Hobey Elite",                 "7021-00",  "OverGarment"),
    ("Gal Suit",                    "7004-01",  "OverGarment"),
    ("Sniper Assault",              "7033-00",  "OverGarment"),
    ("Poncho SRV",                  "7018-00",  "OverGarment"),
    ("Poncho Inbar",                "7030-01",  "OverGarment"),
    ("Poncho Sahar",                "7031-01",  "OverGarment"),
    ("Poncho Sahar 3D",             "",         "OverGarment"),
    ("Poncho Inbar 3D",             "",         "OverGarment"),
    ("SWAT 3-Part SUIT",            "9097-00",  "OverGarment"),
    ("May Suit",                    "7040-01",  "OverGarment"),
    ("May Long Shirt",              "",         "OverGarment"),
    ("Arctic Jacket",               "7051-00",  "OverGarment"),
    ("Arctic Overpants",            "7052-00",  "OverGarment"),
    ("Arctic Poncho SRV",           "",         "OverGarment"),
    ("Arctic Gaiters",              "",         "OverGarment"),
    ("Arctic Suit",                 "",         "OverGarment"),
    ("Gaiters",                     "5004-00",  "OverGarment"),
    ("Survival Blanket",            "9021-00",  "Blankets"),
    ("Water Proof Survival Blanket","9028-00",  "Blankets"),
    ("Water Proof Sniper Blanket",  "9062-00",  "Blankets"),
    ("3D Survival Blanket",         "9029-00",  "Blankets"),
    ("Arctic Survival Blanket",     "",         "Blankets"),
    ("IPS",                         "3007-00",  "Blankets"),
    ("Survival Blanket: Loki",      "",         "Blankets"),
    ("\u05d9\u05e8\u05d9\u05e2\u05ea \u05d4\u05d7\u05e9\u05db\u05d4","","Blankets"),
    ("Operators HS",                "9001-00",  "Hide Site"),
    ("OverWatch (ATGM) HS",         "9004-01",  "Hide Site"),
    ("6X6 Hide Site",               "9095-00",  "Hide Site"),
    ("\u05e2\u05de\u05d3\u05ea \u05d7\u05d5\u05d3","","Hide Site"),
    ("Window Kit",                  "9083-00",  "Urban"),
    ("Balaclava",                   "8006-00",  "Uniform"),
    ("Gaiters",                     "5004-00",  "Uniform"),
    ("Headcover",                   "8030-00",  "Uniform"),
    ("Boonie hat",                  "7022-00",  "Uniform"),
    ("Vest cover",                  "7017-00",  "Uniform"),
    ("Pants",                       "",         "Uniform"),
    ("Shirt",                       "",         "Uniform"),
    ("LC79",                        "9026-00",  "Platform Hide Site"),
    ("Vehicle hide site 9*12",      "9090-00",  "Platform Hide Site"),
    ("Vehicle hide site 9*13",      "9093-00",  "Platform Hide Site"),
    ("CMM 7.5*7.5",                 "3052-00",  "Platform Hide Site"),
    ("CMM 10*10",                   "3077-05",  "Platform Hide Site"),
    ("Zodiac Shore Cover",          "3005-00",  "Platform Hide Site"),
    ("MRZR cover",                  "3003-00",  "Platform Hide Site"),
    ("\u05e8\u05e9\u05ea \u05d4\u05e6\u05dc\u05dc\u05d4","","Platform Hide Site"),
    ("GMV",                         "",         "Platform Hide Site"),
    ("Blackout hood",               "9089-00",  "Accessories"),
    ("Thermal drone landing pad",   "9077-00",  "Accessories"),
    ("Bag cover 90L",               "7009-01",  "Accessories"),
    ("Weapon Wraps",                "5001-00",  "Accessories"),
    ("Loki",                        "",         "Platform On-The-Move"),
    ("HS Floor",                    "",         "Accessories"),
    ("\u05d7\u05dc\u05e7\u05d9\u05dd \u05e9\u05dc \u05e2\u05de\u05d3\u05d5\u05ea","","Accessories"),
]  # 53 → rows 2-54

FABRICS  = ["הכל","RIPSTOP","Mesh","RIPSTOP Double Layer","3D","LOKI Material",
             "WaterProof Blackout Fabric","Beti","MRG","\u05e8\u05e9\u05ea \u05e8\u05db\u05d1",
             "SMT","Meron","Arber","IRR","Stretch","Durable","\u05d0\u05d7\u05e8"]  # 17
STATUSES = ["\u05d4\u05db\u05dc","\u05d1\u05de\u05dc\u05d0\u05d9",
             "\u05d1\u05d7\u05d3\u05e8 \u05ea\u05e6\u05d5\u05d2\u05d4",
             "\u05de\u05d5\u05e9\u05d0\u05dc",
             "\u05d1\u05ea\u05d9\u05e7 \u05d4\u05d3\u05d2\u05de\u05d4"]  # 5
YES_NO   = ["\u05d4\u05db\u05dc","\u05db\u05df","\u05dc\u05d0"]  # 3
SIZES    = ["\u05d4\u05db\u05dc","XS","S","M","L","XL","XXL","XXXL"]  # 8

# ── STYLE HELPERS ─────────────────────────────────────────────────────
def fill(hex_col):
    return PatternFill("solid", fgColor=hex_col)

def hdr_font(color="FFFFFF", bold=True, size=10):
    return Font(bold=bold, color=color, size=size)

def center(wrap=False):
    return Alignment(horizontal="center", vertical="center", wrap_text=wrap)

C_DARK = "1B4332"
C_MID  = "2D6A4F"
C_LITE = "D8F3DC"
C_GOLD = "B7950B"

# ── WORKBOOK ──────────────────────────────────────────────────────────
wb = Workbook()
wb.remove(wb.active)

ws_inv  = wb.create_sheet("\u05de\u05dc\u05d0\u05d9")                   # מלאי
ws_set  = wb.create_sheet("\u05d4\u05d2\u05d3\u05e8\u05d5\u05ea")        # הגדרות
ws_loan = wb.create_sheet("\u05d4\u05e9\u05d0\u05dc\u05d5\u05ea")        # השאלות
ws_tech = wb.create_sheet("\u05e2\u05e8\u05db\u05d9\u05dd \u05d8\u05db\u05e0\u05d5\u05dc\u05d5\u05d2\u05d9\u05d9\u05dd")  # ערכים טכנולוגיים
ws_cont = wb.create_sheet("\u05ea\u05db\u05d5\u05dc\u05d4")              # תכולה
ws_dash = wb.create_sheet("\u05dc\u05d5\u05d7 \u05d1\u05e7\u05e8\u05d4") # לוח בקרה
ws_tops = wb.create_sheet("\u05d8\u05db\u05e0\u05d5")                    # טכנו

for ws in (ws_inv, ws_set, ws_loan, ws_tech, ws_cont, ws_dash, ws_tops):
    ws.sheet_view.rightToLeft = True

print("Sheets created.")

# ═══════════════════════════════════════════════════════════════════════
# PHASE B — SETTINGS SHEET
# ═══════════════════════════════════════════════════════════════════════
ws = ws_set

# ── Prints table A1:D18 ─────────────────────────────────────────────
for c, h in enumerate(["\u05d4\u05d3\u05e4\u05e1",
                        "\u05e1\u05d1\u05d9\u05d1\u05d4 \u05e8\u05d0\u05e9\u05d9\u05ea",
                        "\u05e1\u05d1\u05d9\u05d1\u05d4 \u05de\u05e9\u05e0\u05d9\u05ea",
                        "\u05ea\u05de\u05d5\u05e0\u05d4"], start=1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = hdr_font(); cell.fill = fill(C_DARK); cell.alignment = center()
for i, (name, e1, e2) in enumerate(PRINTS, start=2):
    ws.cell(row=i, column=1, value=name)
    ws.cell(row=i, column=2, value=e1)
    ws.cell(row=i, column=3, value=e2)

# ── Environments F1:F8 ──────────────────────────────────────────────
ws.cell(row=1, column=6, value="\u05e1\u05d1\u05d9\u05d1\u05d5\u05ea").font = hdr_font()
ws.cell(row=1, column=6).fill = fill(C_MID); ws.cell(row=1, column=6).alignment = center()
for i, e in enumerate(ENVIRONMENTS, start=2):
    ws.cell(row=i, column=6, value=e)

# ── Product lines G1:G10 ────────────────────────────────────────────
ws.cell(row=1, column=7, value="\u05e7\u05d5\u05d5\u05d9 \u05de\u05d5\u05e6\u05e8").font = hdr_font()
ws.cell(row=1, column=7).fill = fill(C_MID); ws.cell(row=1, column=7).alignment = center()
for i, ln in enumerate(PRODUCT_LINES, start=2):
    ws.cell(row=i, column=7, value=ln)

# ── Products H1:J54 ─────────────────────────────────────────────────
for c, h in enumerate(["\u05de\u05d5\u05e6\u05e8", '\u05de\u05e7"\u05d8', "\u05e7\u05d5 \u05de\u05d5\u05e6\u05e8"], start=8):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = hdr_font(); cell.fill = fill(C_DARK); cell.alignment = center()
for i, (name, sku, line) in enumerate(PRODUCTS, start=2):
    ws.cell(row=i, column=8, value=name)
    ws.cell(row=i, column=9, value=sku)
    ws.cell(row=i, column=10, value=line)

# ── Fabrics L1:L18 ──────────────────────────────────────────────────
ws.cell(row=1, column=12, value="\u05e1\u05d5\u05d2\u05d9 \u05d1\u05d3").font = hdr_font()
ws.cell(row=1, column=12).fill = fill(C_MID); ws.cell(row=1, column=12).alignment = center()
for i, f in enumerate(FABRICS, start=2):
    ws.cell(row=i, column=12, value=f)

# ── Print names list M1:M19  (הכל + 17 prints) ──────────────────────
ws.cell(row=1, column=13, value="\u05d4\u05d3\u05e4\u05e1\u05d9\u05dd").font = hdr_font()
ws.cell(row=1, column=13).fill = fill(C_MID); ws.cell(row=1, column=13).alignment = center()
ws.cell(row=2, column=13, value="\u05d4\u05db\u05dc")
for i, (name, _, _) in enumerate(PRINTS, start=3):
    ws.cell(row=i, column=13, value=name)

# ── Statuses N1:N6 ──────────────────────────────────────────────────
ws.cell(row=1, column=14, value="\u05e1\u05d8\u05d8\u05d5\u05e1\u05d9\u05dd").font = hdr_font()
ws.cell(row=1, column=14).fill = fill(C_DARK); ws.cell(row=1, column=14).alignment = center()
for i, s in enumerate(STATUSES, start=2):
    ws.cell(row=i, column=14, value=s)

# ── Yes/No O1:O4 ────────────────────────────────────────────────────
ws.cell(row=1, column=15, value="\u05db\u05df/\u05dc\u05d0").font = hdr_font()
ws.cell(row=1, column=15).fill = fill(C_MID); ws.cell(row=1, column=15).alignment = center()
for i, yn in enumerate(YES_NO, start=2):
    ws.cell(row=i, column=15, value=yn)

# ── Sizes P1:P9 ─────────────────────────────────────────────────────
ws.cell(row=1, column=16, value="\u05de\u05d9\u05d3\u05d5\u05ea").font = hdr_font()
ws.cell(row=1, column=16).fill = fill(C_MID); ws.cell(row=1, column=16).alignment = center()
for i, sz in enumerate(SIZES, start=2):
    ws.cell(row=i, column=16, value=sz)

# ── Contents reference R1:S1 ────────────────────────────────────────
ws.cell(row=1, column=18, value="\u05de\u05d5\u05e6\u05e8 \u05ea\u05db\u05d5\u05dc\u05d4").font = hdr_font()
ws.cell(row=1, column=18).fill = fill(C_DARK); ws.cell(row=1, column=18).alignment = center()
ws.cell(row=1, column=19, value="\u05ea\u05db\u05d5\u05dc\u05d4 \u05e1\u05d8\u05e0\u05d3\u05e8\u05d8\u05d9\u05ea").font = hdr_font()
ws.cell(row=1, column=19).fill = fill(C_DARK); ws.cell(row=1, column=19).alignment = center()

# Col widths
for c, w in {1:18,2:16,3:20,4:10,6:20,7:22,8:30,9:12,10:22,12:24,13:20,14:20,15:10,16:10,18:24,19:30}.items():
    ws.column_dimensions[get_column_letter(c)].width = w

print("Settings sheet built.")

# ═══════════════════════════════════════════════════════════════════════
# NAMED RANGES
# ═══════════════════════════════════════════════════════════════════════
# Sheet name quoting: Hebrew name needs single quotes in range refs
SN = "'\u05d4\u05d2\u05d3\u05e8\u05d5\u05ea'"  # 'הגדרות'

def add_nr(name, ref):
    dn = DefinedName(name=name, attr_text=ref)
    wb.defined_names[name] = dn

add_nr("PrintsTable",  f"{SN}!$A$2:$D$18")
add_nr("EnvList",      f"{SN}!$F$2:$F$8")
add_nr("LineList",     f"{SN}!$G$2:$G$10")
add_nr("ProductsTable",f"{SN}!$H$2:$J$54")
add_nr("ProductsName", f"{SN}!$H$2:$H$54")
add_nr("ProductsSKU",  f"{SN}!$I$2:$I$54")
add_nr("FabricList",   f"{SN}!$L$2:$L$18")
add_nr("PrintsList",   f"{SN}!$M$2:$M$19")
add_nr("StatusList",   f"{SN}!$N$2:$N$6")
add_nr("YesNoList",    f"{SN}!$O$2:$O$4")
add_nr("SizesList",    f"{SN}!$P$2:$P$9")

print("Named ranges added.")


# ═══════════════════════════════════════════════════════════════════════
# PHASE A — INVENTORY SHEET
# ═══════════════════════════════════════════════════════════════════════
ws = ws_inv
MAX_ROW = 357

# Sheet refs for formulas
SET_S  = "'\u05d4\u05d2\u05d3\u05e8\u05d5\u05ea'"        # 'הגדרות'
DASH_S = "'\u05dc\u05d5\u05d7 \u05d1\u05e7\u05e8\u05d4'" # 'לוח בקרה'
TECH_S = "'\u05d8\u05db\u05e0\u05d5'"                     # 'טכנו'

# ── Row 1: category header row (merged) ─────────────────────────────
def cat_cell(ws, col_start, col_end, row, label, fg):
    if col_start != col_end:
        ws.merge_cells(start_row=row, start_column=col_start,
                       end_row=row,   end_column=col_end)
    c = ws.cell(row=row, column=col_start, value=label)
    c.font = Font(bold=True, color="FFFFFF", size=10)
    c.fill = fill(fg); c.alignment = center(wrap=True)

cat_cell(ws,  1,  5, 1, "\u05d6\u05d9\u05d4\u05d5\u05d9 \u05de\u05d5\u05e6\u05e8",   C_DARK)  # זיהוי מוצר
cat_cell(ws,  6,  7, 1, "\u05d1\u05d3",                                                 C_MID)   # בד
cat_cell(ws,  8,  9, 1, "\u05d4\u05d3\u05e4\u05e1\u05d9\u05dd",                        C_MID)   # הדפסים
cat_cell(ws, 10, 15, 1, "\u05ea\u05e7\u05d9\u05e0\u05d5\u05ea \u05d5\u05de\u05d0\u05e4\u05d9\u05d9\u05e0\u05d9\u05dd", C_DARK) # תקינות ומאפיינים
cat_cell(ws, 16, 17, 1, "\u05d4\u05e2\u05e8\u05d5\u05ea",                              C_MID)   # הערות
# R1, S1 empty  → columns 18-19
cat_cell(ws, 20, 20, 1, "\u05e1\u05d8\u05d8\u05d5\u05e1",                              C_DARK)  # סטטוס  T
cat_cell(ws, 21, 21, 1, "\u05d4\u05e2\u05e8\u05d5\u05ea",                              C_DARK)  # הערות  U
cat_cell(ws, 22, 25, 1, "\u05d4\u05e9\u05d0\u05dc\u05d4",                              C_MID)   # השאלה  V-Y

# ── Row 2: column headers ────────────────────────────────────────────
ROW2 = [
    "\u05e7\u05d5 \u05de\u05d5\u05e6\u05e8",            # A  קו מוצר
    "\u05e9\u05dd \u05de\u05d5\u05e6\u05e8",             # B  שם מוצר
    "\u05d2\u05e8\u05e1\u05d4",                           # C  גרסה
    "\u05de\u05d6\u05d4\u05d4 \u05d9\u05d9\u05d7\u05d5\u05d3\u05d9", # D מזהה ייחודי
    '\u05de\u05e7"\u05d8 (\u05d0\u05d5\u05d8\u05d5\u05de\u05d8\u05d9)', # E מק"ט
    "\u05e1\u05d5\u05d2 \u05d1\u05d3 1",                 # F  סוג בד 1
    "\u05e1\u05d5\u05d2 \u05d1\u05d3 2",                 # G  סוג בד 2
    "\u05d4\u05d3\u05e4\u05e1 \u05e6\u05d3 A'",          # H  הדפס צד A'
    "\u05d4\u05d3\u05e4\u05e1 \u05e6\u05d3 \u05d1'",     # I  הדפס צד ב'
    "\u05ea\u05e7\u05d9\u05df \u05d5\u05d9\u05d6\u05d5\u05d0\u05dc\u05d9\u05ea \u05d0'", # J
    "\u05ea\u05e7\u05d9\u05df \u05d5\u05d9\u05d6\u05d5\u05d0\u05dc\u05d9\u05ea \u05d1'", # K
    "\u05ea\u05e7\u05d9\u05df \u05ea\u05e8\u05de\u05d9\u05ea \u05e6\u05d3 \u05d0'",       # L
    "\u05ea\u05e7\u05d9\u05df \u05ea\u05e8\u05de\u05d9\u05ea \u05e6\u05d3 \u05d1'",       # M
    "\u05e4\u05d9\u05e6'\u05e8\u05d9\u05dd \u05ea\u05e7\u05d9\u05e0\u05d9\u05dd",         # N
    "\u05d2\u05e8\u05e1\u05d4 \u05e1\u05d5\u05e4\u05d9\u05ea",                           # O
    "\u05d4\u05e2\u05e8\u05d5\u05ea \u05d2\u05e8\u05e1\u05d4/\u05e4\u05d9\u05e6'\u05e8\u05d9\u05dd", # P
    "\u05d2\u05e8\u05e1\u05d4 \u05de\u05d9\u05d5\u05d7\u05d3\u05ea(\u05e4\u05d9\u05ea\u05d5\u05d7)", # Q
    "\u05de\u05d9\u05d3\u05d4",                           # R  מידה
    "\u05e4\u05e2\u05e8 \u05d1\u05ea\u05db\u05d5\u05dc\u05d4",                          # S  פער בתכולה
    "\u05e1\u05d8\u05d8\u05d5\u05e1",                     # T  סטטוס
    "\u05d4\u05e2\u05e8\u05d5\u05ea",                     # U  הערות
    "\u05de\u05d5\u05e9\u05d0\u05dc \u05dc",              # V  מושאל ל
    "\u05de\u05d3\u05d9\u05e0\u05d4",                     # W  מדינה
    "\u05d0\u05d7\u05e8\u05d0\u05d9 \u05de\u05d4\u05d7\u05d1\u05e8\u05d4",             # X  אחראי מהחברה
    "\u05d4\u05d7\u05d6\u05e8\u05d4 \u05de\u05e9\u05d5\u05e2\u05e8\u05ea",             # Y  החזרה משוערת
    "\u05e1\u05d1\u05d9\u05d1\u05d4 A'",                  # Z
    "\u05e1\u05d1\u05d9\u05d1\u05d4 \u05de\u05e9\u05e0\u05d9\u05ea A'",               # AA
    "\u05e1\u05d1\u05d9\u05d1\u05d4 \u05d1'",             # AB
    "\u05e1\u05d1\u05d9\u05d1\u05d4 \u05de\u05e9\u05e0\u05d9\u05ea \u05d1'",          # AC
    "\u05de\u05e1\u05e0\u05df \u05d1\u05e7\u05e8\u05d4",  # AD
    "\u05de\u05e1\u05e0\u05df \u05d8\u05db\u05e0\u05d5",  # AE
]

for col_idx, header in enumerate(ROW2, start=1):
    c = ws.cell(row=2, column=col_idx, value=header)
    c.font = Font(bold=True, color="FFFFFF", size=9)
    c.fill = fill(C_MID); c.alignment = center(wrap=True)

ws.row_dimensions[1].height = 22
ws.row_dimensions[2].height = 36

# ── Column widths ───────────────────────────────────────────────────
widths = {1:14,2:22,3:12,4:15,5:11,6:14,7:14,8:15,9:15,10:8,11:8,12:8,
          13:8,14:8,15:8,16:20,17:12,18:7,19:7,20:14,21:18,22:15,23:11,
          24:15,25:13,26:11,27:14,28:11,29:14,30:4,31:4}
for col, w in widths.items():
    ws.column_dimensions[get_column_letter(col)].width = w

# Hide formula/helper columns
for col_letter in ("E","Z","AA","AB","AC","AD","AE"):
    ws.column_dimensions[col_letter].hidden = True

ws.freeze_panes = "A3"

# ── Data rows: formulas 3-357 ────────────────────────────────────────
for r in range(3, MAX_ROW + 1):
    # E: SKU lookup from settings
    ws.cell(row=r, column=5).value = (
        f'=IFERROR(INDEX({SET_S}!$I:$I,MATCH(B{r},{SET_S}!$H:$H,0)),"")'
    )
    # Z: primary env for print A
    ws.cell(row=r, column=26).value = (
        f'=IFERROR(VLOOKUP(H{r},{SET_S}!$A:$D,2,FALSE),"")'
    )
    # AA: secondary env for print A
    ws.cell(row=r, column=27).value = (
        f'=IFERROR(VLOOKUP(H{r},{SET_S}!$A:$D,3,FALSE),"")'
    )
    # AB: primary env for print B
    ws.cell(row=r, column=28).value = (
        f'=IFERROR(VLOOKUP(I{r},{SET_S}!$A:$D,2,FALSE),"")'
    )
    # AC: secondary env for print B
    ws.cell(row=r, column=29).value = (
        f'=IFERROR(VLOOKUP(I{r},{SET_S}!$A:$D,3,FALSE),"")'
    )
    # AD: dashboard filter helper — returns relative row index if row matches ALL filters
    ws.cell(row=r, column=30).value = (
        f'=IF(AND('
        f'OR({DASH_S}!$C$3="\u05d4\u05db\u05dc",Z{r}={DASH_S}!$C$3,AA{r}={DASH_S}!$C$3,AB{r}={DASH_S}!$C$3,AC{r}={DASH_S}!$C$3),'
        f'OR({DASH_S}!$C$4="\u05d4\u05db\u05dc",A{r}={DASH_S}!$C$4),'
        f'OR({DASH_S}!$C$5="\u05d4\u05db\u05dc",T{r}={DASH_S}!$C$5),'
        f'OR({DASH_S}!$C$6="\u05d4\u05db\u05dc",H{r}={DASH_S}!$C$6,I{r}={DASH_S}!$C$6),'
        f'OR({DASH_S}!$C$7="\u05d4\u05db\u05dc",J{r}={DASH_S}!$C$7),'
        f'OR({DASH_S}!$C$8="\u05d4\u05db\u05dc",L{r}={DASH_S}!$C$8),'
        f'OR({DASH_S}!$C$9="\u05d4\u05db\u05dc",Q{r}={DASH_S}!$C$9),'
        f'OR({DASH_S}!$C$10="\u05d4\u05db\u05dc",F{r}={DASH_S}!$C$10,G{r}={DASH_S}!$C$10),'
        f'D{r}<>""'
        f'),ROW()-2,"")'
    )
    # AE: techops filter helper
    ws.cell(row=r, column=31).value = (
        f'=IF(AND('
        f'OR({TECH_S}!$C$3="\u05d4\u05db\u05dc",Z{r}={TECH_S}!$C$3,AA{r}={TECH_S}!$C$3,AB{r}={TECH_S}!$C$3,AC{r}={TECH_S}!$C$3),'
        f'OR({TECH_S}!$C$4="\u05d4\u05db\u05dc",A{r}={TECH_S}!$C$4),'
        f'D{r}<>""'
        f'),ROW()-2,"")'
    )

# ── Data validations ─────────────────────────────────────────────────
def add_dv(ws, formula1, sqref, allow_blank=True):
    dv = DataValidation(type="list", formula1=formula1,
                        allow_blank=allow_blank, showDropDown=False)
    dv.sqref = sqref
    ws.add_data_validation(dv)

SR  = f"{SN}!$G$2:$G$10"   # LineList range
FR  = f"{SN}!$L$2:$L$18"   # FabricList
PR  = f"{SN}!$M$2:$M$19"   # PrintsList
STR = f"{SN}!$N$2:$N$6"    # StatusList
YNR = f"{SN}!$O$2:$O$4"    # YesNoList
SZR = f"{SN}!$P$2:$P$9"    # SizesList

add_dv(ws_inv, SR,  f"A3:A{MAX_ROW}")
add_dv(ws_inv, FR,  f"F3:G{MAX_ROW}")
add_dv(ws_inv, PR,  f"H3:I{MAX_ROW}")
add_dv(ws_inv, YNR, f"J3:O{MAX_ROW}")
add_dv(ws_inv, YNR, f"Q3:Q{MAX_ROW}")
add_dv(ws_inv, YNR, f"S3:S{MAX_ROW}")
add_dv(ws_inv, STR, f"T3:T{MAX_ROW}")
add_dv(ws_inv, SZR, f"R3:R{MAX_ROW}")

print("Inventory sheet built.")


# ═══════════════════════════════════════════════════════════════════════
# PHASE D — LOANS SHEET
# ═══════════════════════════════════════════════════════════════════════
ws = ws_loan

# Title
t = ws.cell(row=1, column=1, value="\u05d4\u05e9\u05d0\u05dc\u05d5\u05ea \u05e4\u05e2\u05d9\u05dc\u05d5\u05ea \u2013 \u05de\u05d5\u05e9\u05d0\u05dc\u05d9\u05dd \u05e0\u05d5\u05db\u05d7\u05d9\u05d9\u05dd")  # השאלות פעילות
t.font = Font(bold=True, size=14, color=C_DARK); t.alignment = center()
ws.merge_cells("A1:M1")

ws.cell(row=3, column=1,
        value="\u05d4\u05d8\u05d1\u05dc\u05d4 \u05e0\u05d8\u05e2\u05e0\u05ea \u05d0\u05d5\u05d8\u05d5\u05de\u05d8\u05d9\u05ea \u05de\u05d2\u05d9\u05dc\u05d9\u05d5\u05df \u05d4\u05de\u05dc\u05d0\u05d9 (\u05e1\u05d8\u05d8\u05d5\u05e1=\u05de\u05d5\u05e9\u05d0\u05dc). \u05e2\u05d3\u05db\u05df \u05e1\u05d8\u05d8\u05d5\u05e1 \u05d1\u05d2\u05d9\u05dc\u05d9\u05d5\u05df \u05d4\u05de\u05dc\u05d0\u05d9 \u05db\u05d3\u05d9 \u05dc\u05e8\u05e2\u05e0\u05df.")
ws.cell(row=3, column=1).font = Font(italic=True, size=9, color="666666")
ws.merge_cells("A3:M3")

# Row 4: headers
LOAN_HDRS = [
    '\u05de\u05e1\'',               # מס'  A
    '\u05de\u05d6\u05d4\u05d4',     # מזהה B
    '\u05e9\u05dd \u05de\u05d5\u05e6\u05e8',  # שם מוצר C
    '\u05e7\u05d5 \u05de\u05d5\u05e6\u05e8',  # קו מוצר D
    '\u05de\u05d5\u05e9\u05d0\u05dc \u05dc',  # מושאל ל E
    '\u05de\u05d3\u05d9\u05e0\u05d4',         # מדינה F
    '\u05d0\u05d7\u05e8\u05d0\u05d9',         # אחראי G
    '\u05ea\u05d0\u05e8\u05d9\u05da \u05d4\u05e9\u05d0\u05dc\u05d4',  # תאריך השאלה H
    '\u05d4\u05d7\u05d6\u05e8\u05d4 \u05de\u05e9\u05d5\u05e2\u05e8\u05ea', # החזרה משוערת I
    '\u05d4\u05d7\u05d6\u05e8\u05d4 \u05d1\u05e4\u05d5\u05e2\u05dc',  # החזרה בפועל J
    '\u05e1\u05d8\u05d8\u05d5\u05e1',  # סטטוס K
    '\u05d4\u05ea\u05e8\u05d0\u05d4',  # התראה L
    '\u05d4\u05e2\u05e8\u05d5\u05ea',  # הערות M
]
for col_idx, h in enumerate(LOAN_HDRS, start=1):
    c = ws.cell(row=4, column=col_idx, value=h)
    c.font = hdr_font(); c.fill = fill(C_DARK); c.alignment = center(wrap=True)
ws.row_dimensions[4].height = 28

# Rows 5..100: dynamic formulas pulling from inventory מלאי
INV_S = "'\u05de\u05dc\u05d0\u05d9'"  # 'מלאי'
MAX_LOAN = 100
for r in range(5, MAX_LOAN + 1):
    k = r - 4  # rank (1,2,3,...)
    # A: row number
    ws.cell(row=r, column=1).value = f'=IF(B{r}="","",{r-4})'
    # B: מזהה — AGGREGATE(15) = SMALL, option 6 = ignore errors → finds nth item where T=מושאל
    ws.cell(row=r, column=2).value = (
        f'=IFERROR(INDEX({INV_S}!$D$3:$D$357,'
        f'AGGREGATE(15,6,(ROW({INV_S}!$D$3:$D$357)-ROW({INV_S}!$D$3)+1)'
        f'/({INV_S}!$T$3:$T$357="\u05de\u05d5\u05e9\u05d0\u05dc"),{k})),"")'
    )
    # C-G: auto-populated from inventory by ID
    ws.cell(row=r, column=3).value  = f'=IF(B{r}="","",IFERROR(INDEX({INV_S}!$B$3:$B$357,MATCH(B{r},{INV_S}!$D$3:$D$357,0)),""))'
    ws.cell(row=r, column=4).value  = f'=IF(B{r}="","",IFERROR(INDEX({INV_S}!$A$3:$A$357,MATCH(B{r},{INV_S}!$D$3:$D$357,0)),""))'
    ws.cell(row=r, column=5).value  = f'=IF(B{r}="","",IFERROR(INDEX({INV_S}!$V$3:$V$357,MATCH(B{r},{INV_S}!$D$3:$D$357,0)),""))'
    ws.cell(row=r, column=6).value  = f'=IF(B{r}="","",IFERROR(INDEX({INV_S}!$W$3:$W$357,MATCH(B{r},{INV_S}!$D$3:$D$357,0)),""))'
    ws.cell(row=r, column=7).value  = f'=IF(B{r}="","",IFERROR(INDEX({INV_S}!$X$3:$X$357,MATCH(B{r},{INV_S}!$D$3:$D$357,0)),""))'
    # H: loan date — manual entry (left blank)
    # I: expected return — from inventory
    ws.cell(row=r, column=9).value  = f'=IF(B{r}="","",IFERROR(INDEX({INV_S}!$Y$3:$Y$357,MATCH(B{r},{INV_S}!$D$3:$D$357,0)),""))'
    # J: actual return — manual entry (blank)
    # K: status
    ws.cell(row=r, column=11).value = (
        f'=IF(B{r}="","",IF(J{r}<>"","\u05d4\u05d5\u05d7\u05d6\u05e8","\u05e4\u05e2\u05d9\u05dc"))'
    )
    # L: alert
    ws.cell(row=r, column=12).value = (
        f'=IF(AND(K{r}="\u05e4\u05e2\u05d9\u05dc",I{r}<>""),'
        f'"\u05e2\u05d3\u05d9\u05d9\u05df \u05de\u05d5\u05e9\u05d0\u05dc","")'
    )
    # M: notes — manual

# Column widths
for c, w in {1:5,2:15,3:22,4:18,5:18,6:12,7:16,8:14,9:14,10:14,11:12,12:14,13:18}.items():
    ws.column_dimensions[get_column_letter(c)].width = w
ws.freeze_panes = "A5"

print("Loans sheet built.")


# ═══════════════════════════════════════════════════════════════════════
# PHASE C — TECH VALUES SHEET
# ═══════════════════════════════════════════════════════════════════════
ws = ws_tech

# Title
t = ws.cell(row=1, column=1, value="\u05e2\u05e8\u05db\u05d9\u05dd \u05d8\u05db\u05e0\u05d5\u05dc\u05d5\u05d2\u05d9\u05d9\u05dd")
t.font = Font(bold=True, size=14, color=C_DARK)
ws.merge_cells("A1:O1")

ws.cell(row=3, column=1,
        value="\u05e2\u05d3\u05db\u05df \u05de\u05d6\u05d4\u05d4 (\u05e2\u05de\u05d5\u05d3\u05d4 A) \u05db\u05d3\u05d9 \u05dc\u05d4\u05e6\u05d9\u05d2 \u05e2\u05e8\u05db\u05d9\u05dd \u05d0\u05d5\u05d8\u05d5\u05de\u05d8\u05d9\u05ea. \u05e9\u05d3\u05d5\u05ea MWIR/LWIR/NIR \u05d5\u05e0\u05d5\u05d1\u05dc \u05d4\u05dd \u05d4\u05d6\u05e0\u05d4 \u05d9\u05d3\u05e0\u05d9\u05ea.")
ws.cell(row=3, column=1).font = Font(italic=True, size=9, color="666666")
ws.merge_cells("A3:O3")

# Row 4: headers
TECH_HDRS = [
    "\u05de\u05d6\u05d4\u05d4",                     # A מזהה
    "\u05e9\u05dd \u05de\u05d5\u05e6\u05e8",         # B שם מוצר
    "\u05e7\u05d5 \u05de\u05d5\u05e6\u05e8",         # C קו מוצר
    "\u05d4\u05d3\u05e4\u05e1 \u05e6\u05d3 \u05d0'", # D הדפס צד א'
    "MWIR \u05d0'",                                   # E
    "LWIR \u05d0'",                                   # F
    "NIR \u05d0'",                                    # G
    "\u05d4\u05d3\u05e4\u05e1 \u05e6\u05d3 \u05d1'", # H הדפס צד ב'
    "MWIR \u05d1'",                                   # I
    "LWIR \u05d1'",                                   # J
    "NIR \u05d1'",                                    # K
    "\u05e0\u05d5\u05d1\u05dc",                       # L נובל
    "\u05de\u05de\u05d5\u05e6\u05e2 MWIR",            # M ממוצע MWIR
    "\u05de\u05de\u05d5\u05e6\u05e2 LWIR",            # N ממוצע LWIR
    "\u05de\u05de\u05d5\u05e6\u05e2 \u05de\u05e9\u05d5\u05dc\u05d1",  # O ממוצע משולב
    # P: hidden filter helper
    "\u05de\u05e1\u05e0\u05df \u05d8\u05db\u05e0\u05d5",  # P hidden
]
for col_idx, h in enumerate(TECH_HDRS, start=1):
    c = ws.cell(row=4, column=col_idx, value=h)
    c.font = hdr_font(); c.fill = fill(C_DARK); c.alignment = center(wrap=True)
ws.row_dimensions[4].height = 28

# Hide P column
ws.column_dimensions["P"].hidden = True

# Rows 5..357: auto-formulas
for r in range(5, MAX_ROW + 1):
    # B: product name from inventory by ID
    ws.cell(row=r, column=2).value  = f'=IF(A{r}="","",IFERROR(INDEX({INV_S}!$B$3:$B$357,MATCH(A{r},{INV_S}!$D$3:$D$357,0)),""))'
    # C: product line
    ws.cell(row=r, column=3).value  = f'=IF(A{r}="","",IFERROR(INDEX({INV_S}!$A$3:$A$357,MATCH(A{r},{INV_S}!$D$3:$D$357,0)),""))'
    # D: print A from inventory
    ws.cell(row=r, column=4).value  = f'=IF(A{r}="","",IFERROR(INDEX({INV_S}!$H$3:$H$357,MATCH(A{r},{INV_S}!$D$3:$D$357,0)),""))'
    # H: print B from inventory
    ws.cell(row=r, column=8).value  = f'=IF(A{r}="","",IFERROR(INDEX({INV_S}!$I$3:$I$357,MATCH(A{r},{INV_S}!$D$3:$D$357,0)),""))'
    # E,F,G,I,J,K,L: manual entry → leave blank
    # M: avg MWIR
    ws.cell(row=r, column=13).value = (
        f'=IF(AND(E{r}="",I{r}=""),"",IF(H{r}="",E{r},IF(I{r}="",E{r},'
        f'IFERROR((VALUE(E{r})+VALUE(I{r}))/2,E{r}))))'
    )
    # N: avg LWIR
    ws.cell(row=r, column=14).value = (
        f'=IF(AND(F{r}="",J{r}=""),"",IF(H{r}="",F{r},IF(J{r}="",F{r},'
        f'IFERROR((VALUE(F{r})+VALUE(J{r}))/2,F{r}))))'
    )
    # O: combined avg
    ws.cell(row=r, column=15).value = (
        f'=IF(AND(M{r}="",N{r}=""),"",IF(M{r}="",N{r},IF(N{r}="",M{r},'
        f'IFERROR((VALUE(M{r})+VALUE(N{r}))/2,""))))'
    )
    # P: techops filter helper
    ws.cell(row=r, column=16).value = (
        f'=IF(AND('
        f'OR({TECH_S}!$C$3="\u05d4\u05db\u05dc",'
        f'IFERROR(VLOOKUP(D{r},{SET_S}!$A$2:$D$18,2,FALSE),"")={TECH_S}!$C$3,'
        f'IFERROR(VLOOKUP(D{r},{SET_S}!$A$2:$D$18,3,FALSE),"")={TECH_S}!$C$3),'
        f'OR({TECH_S}!$C$4="\u05d4\u05db\u05dc",C{r}={TECH_S}!$C$4),'
        f'OR({TECH_S}!$C$5="\u05d4\u05db\u05dc",L{r}={TECH_S}!$C$5),'
        f'OR({TECH_S}!$C$6="",IFERROR(VALUE(M{r}),0)>={TECH_S}!$C$6),'
        f'OR({TECH_S}!$C$7="",IFERROR(VALUE(M{r}),0)<={TECH_S}!$C$7),'
        f'A{r}<>""'
        f'),ROW()-4,"")'
    )

# Col widths
for c, w in {1:15,2:24,3:20,4:16,5:9,6:9,7:9,8:16,9:9,10:9,11:9,12:9,13:12,14:12,15:14,16:4}.items():
    ws.column_dimensions[get_column_letter(c)].width = w
ws.freeze_panes = "A5"

# Yes/No validation for Nobel (L)
add_dv(ws_tech, f"{SN}!$O$3:$O$4", f"L5:L{MAX_ROW}")  # just כן/לא (skip הכל)

print("Tech values sheet built.")


# ═══════════════════════════════════════════════════════════════════════
# PHASE G-CONTENTS SHEET (minimal structure)
# ═══════════════════════════════════════════════════════════════════════
ws = ws_cont
ws.cell(row=1, column=1, value="\u05de\u05d5\u05e6\u05e8 \u05ea\u05db\u05d5\u05dc\u05d4").font = hdr_font()
ws.cell(row=1, column=1).fill = fill(C_DARK); ws.cell(row=1, column=1).alignment = center()
ws.cell(row=1, column=2, value="\u05ea\u05db\u05d5\u05dc\u05d4 \u05e1\u05d8\u05e0\u05d3\u05e8\u05d8\u05d9\u05ea").font = hdr_font()
ws.cell(row=1, column=2).fill = fill(C_DARK); ws.cell(row=1, column=2).alignment = center()
ws.column_dimensions["A"].width = 28
ws.column_dimensions["B"].width = 50

print("Contents sheet built.")

# ═══════════════════════════════════════════════════════════════════════
# PHASE E — REGULAR DASHBOARD (לוח בקרה)
# ═══════════════════════════════════════════════════════════════════════
ws = ws_dash

# Title
t = ws.cell(row=1, column=1, value="\u05dc\u05d5\u05d7 \u05d1\u05e7\u05e8\u05d4 \u2013 \u05de\u05dc\u05d0\u05d9 \u05de\u05d3\u05d2\u05d9\u05de\u05d9 \u05e9\u05d9\u05d5\u05d5\u05e7")
t.font = Font(bold=True, size=16, color=C_DARK); t.alignment = center()
ws.merge_cells("A1:J1")
ws.row_dimensions[1].height = 28

# ── Filter section (rows 3-11) ───────────────────────────────────────
ws.cell(row=2, column=2, value="\u05e1\u05d9\u05e0\u05d5\u05e0").font = Font(bold=True, size=11, color=C_MID)

DASH_FILTERS = [
    # (row, label_B, default_C, formula1_for_validation)
    (3,  "\u05e1\u05d1\u05d9\u05d1\u05d4",              "\u05d4\u05db\u05dc", f"{SN}!$F$2:$F$8"),
    (4,  "\u05e7\u05d5 \u05de\u05d5\u05e6\u05e8",       "\u05d4\u05db\u05dc", f"{SN}!$G$2:$G$10"),
    (5,  "\u05e1\u05d8\u05d8\u05d5\u05e1",               "\u05d4\u05db\u05dc", f"{SN}!$N$2:$N$6"),
    (6,  "\u05d4\u05d3\u05e4\u05e1",                     "\u05d4\u05db\u05dc", f"{SN}!$M$2:$M$19"),
    (7,  "\u05ea\u05e7\u05d9\u05e0\u05d5\u05ea \u05d5\u05d9\u05d6\u05d5\u05d0\u05dc\u05d9\u05ea \u05d0'",  "\u05d4\u05db\u05dc", f"{SN}!$O$2:$O$4"),
    (8,  "\u05ea\u05e7\u05d9\u05e0\u05d5\u05ea \u05ea\u05e8\u05de\u05d9\u05ea \u05d0'",  "\u05d4\u05db\u05dc", f"{SN}!$O$2:$O$4"),
    (9,  "\u05d2\u05e8\u05e1\u05d4 \u05de\u05d9\u05d5\u05d7\u05d3\u05ea",               "\u05d4\u05db\u05dc", f"{SN}!$O$2:$O$4"),
    (10, "\u05e1\u05d5\u05d2 \u05d1\u05d3",              "\u05d4\u05db\u05dc", f"{SN}!$L$2:$L$18"),
]

for row, label, default, formula1 in DASH_FILTERS:
    lc = ws.cell(row=row, column=2, value=label)
    lc.font = Font(bold=True, size=10); lc.alignment = Alignment(horizontal="right", vertical="center")
    vc = ws.cell(row=row, column=3, value=default)
    vc.fill = fill(C_LITE); vc.alignment = center()
    vc.font = Font(size=10)
    add_dv(ws, formula1, f"C{row}")

ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 28
ws.column_dimensions["C"].width = 22

# ── Summary section (rows 13-20) ─────────────────────────────────────
ws.cell(row=13, column=2, value="\u05e1\u05d9\u05db\u05d5\u05dd").font = Font(bold=True, size=11, color=C_MID)

SUMMARY = [
    (14, "\u05e1\u05d4\"\u05db \u05de\u05d5\u05e6\u05e8\u05d9\u05dd \u05de\u05e1\u05d5\u05e0\u05e0\u05d9\u05dd",
         '=SUMPRODUCT((\u05de\u05dc\u05d0\u05d9!$AD$3:$AD$357<>"")*1)'),
    (15, "\u05d1\u05de\u05dc\u05d0\u05d9",
         '=SUMPRODUCT((\u05de\u05dc\u05d0\u05d9!$AD$3:$AD$357<>"")*(\u05de\u05dc\u05d0\u05d9!$T$3:$T$357="\u05d1\u05de\u05dc\u05d0\u05d9"))'),
    (16, "\u05d1\u05d7\u05d3\u05e8 \u05ea\u05e6\u05d5\u05d2\u05d4",
         '=SUMPRODUCT((\u05de\u05dc\u05d0\u05d9!$AD$3:$AD$357<>"")*(\u05de\u05dc\u05d0\u05d9!$T$3:$T$357="\u05d1\u05d7\u05d3\u05e8 \u05ea\u05e6\u05d5\u05d2\u05d4"))'),
    (17, "\u05de\u05d5\u05e9\u05d0\u05dc",
         '=SUMPRODUCT((\u05de\u05dc\u05d0\u05d9!$AD$3:$AD$357<>"")*(\u05de\u05dc\u05d0\u05d9!$T$3:$T$357="\u05de\u05d5\u05e9\u05d0\u05dc"))'),
    (18, "\u05d1\u05ea\u05d9\u05e7 \u05d4\u05d3\u05d2\u05de\u05d4",
         '=SUMPRODUCT((\u05de\u05dc\u05d0\u05d9!$AD$3:$AD$357<>"")*(\u05de\u05dc\u05d0\u05d9!$T$3:$T$357="\u05d1\u05ea\u05d9\u05e7 \u05d4\u05d3\u05d2\u05de\u05d4"))'),
    (19, "\u05ea\u05e7\u05d9\u05df \u05d5\u05d9\u05d6\u05d5\u05d0\u05dc\u05d9\u05ea + \u05ea\u05e8\u05de\u05d9\u05ea",
         '=SUMPRODUCT((\u05de\u05dc\u05d0\u05d9!$AD$3:$AD$357<>"")*((\u05de\u05dc\u05d0\u05d9!$J$3:$J$357="\u05db\u05df")+(\u05de\u05dc\u05d0\u05d9!$J$3:$J$357="1")>0)*((\u05de\u05dc\u05d0\u05d9!$L$3:$L$357="\u05db\u05df")+(\u05de\u05dc\u05d0\u05d9!$L$3:$L$357="1")>0))'),
    (20, "\u05d2\u05e8\u05e1\u05d4 \u05de\u05d9\u05d5\u05d7\u05d3\u05ea (\u05e4\u05d9\u05ea\u05d5\u05d7)",
         '=SUMPRODUCT((\u05de\u05dc\u05d0\u05d9!$AD$3:$AD$357<>"")*(\u05de\u05dc\u05d0\u05d9!$Q$3:$Q$357="\u05db\u05df"))'),
]

for row, label, formula in SUMMARY:
    lc = ws.cell(row=row, column=2, value=label)
    lc.font = Font(size=10)
    vc = ws.cell(row=row, column=3)
    # Store SUMPRODUCT formula using Hebrew sheet name — we must embed the Hebrew directly
    vc.value = formula.replace('\u05de\u05dc\u05d0\u05d9', "'\u05de\u05dc\u05d0\u05d9'")
    vc.font = Font(bold=True, size=11, color=C_DARK); vc.alignment = center()

# ── Detail section headers (row 22) ─────────────────────────────────
DETAIL_HDRS = [
    "\u05e7\u05d5 \u05de\u05d5\u05e6\u05e8",  # A קו מוצר
    "\u05e9\u05dd \u05de\u05d5\u05e6\u05e8",   # B שם מוצר
    "\u05d2\u05e8\u05e1\u05d4",                 # C גרסה
    "\u05de\u05d6\u05d4\u05d4",                 # D מזהה
    "\u05d4\u05d3\u05e4\u05e1 \u05d0'",         # E הדפס א'
    "\u05d4\u05d3\u05e4\u05e1 \u05d1'",         # F הדפס ב'
    "\u05e1\u05d1\u05d9\u05d1\u05d4",           # G סביבה
    "\u05e1\u05d8\u05d8\u05d5\u05e1",           # H סטטוס
    "\u05de\u05d5\u05e9\u05d0\u05dc \u05dc",    # I מושאל ל
    "\u05d4\u05e2\u05e8\u05d5\u05ea",           # J הערות
]
for col_idx, h in enumerate(DETAIL_HDRS, start=1):
    c = ws.cell(row=22, column=col_idx, value=h)
    c.font = hdr_font(); c.fill = fill(C_MID); c.alignment = center(wrap=True)
ws.row_dimensions[22].height = 24

# Detail rows 23..150 pulling from inventory via AD helper
# AD contains ROW()-2 (relative index into D3:D357) for matching rows
INV_DATA_COLS = {
    1: "A",  # קו מוצר
    2: "B",  # שם מוצר
    3: "C",  # גרסה
    4: "D",  # מזהה
    5: "H",  # הדפס א'
    6: "I",  # הדפס ב'
    7: "Z",  # סביבה ראשית א'
    8: "T",  # סטטוס
    9: "V",  # מושאל ל
    10:"U",  # הערות
}

for dr in range(23, 151):
    rank = dr - 22  # 1,2,3,...
    for dc, inv_col in INV_DATA_COLS.items():
        ws.cell(row=dr, column=dc).value = (
            f'=IFERROR(INDEX({INV_S}!${inv_col}$3:${inv_col}$357,'
            f'SMALL({INV_S}!$AD$3:$AD$357,{rank})),"")'
        )

# Col widths for dashboard
for c, w in {1:16,2:22,3:12,4:14,5:14,6:14,7:14,8:14,9:16,10:22}.items():
    ws.column_dimensions[get_column_letter(c)].width = max(
        ws.column_dimensions[get_column_letter(c)].width or 0, w)

ws.freeze_panes = "A23"
print("Dashboard sheet built.")


# ═══════════════════════════════════════════════════════════════════════
# PHASE F — TECHOPS DASHBOARD (טכנו)
# ═══════════════════════════════════════════════════════════════════════
ws = ws_tops

t = ws.cell(row=1, column=1, value="\u05dc\u05d5\u05d7 \u05d8\u05db\u05e0\u05d5-\u05de\u05d1\u05e6\u05e2\u05d9")
t.font = Font(bold=True, size=16, color=C_DARK); t.alignment = center()
ws.merge_cells("A1:J1")
ws.row_dimensions[1].height = 28

ws.cell(row=2, column=2, value="\u05e1\u05d9\u05e0\u05d5\u05e0").font = Font(bold=True, size=11, color=C_MID)

TOPS_FILTERS = [
    (3,  "\u05e1\u05d1\u05d9\u05d1\u05d4",              "\u05d4\u05db\u05dc", f"{SN}!$F$2:$F$8"),
    (4,  "\u05e7\u05d5 \u05de\u05d5\u05e6\u05e8",       "\u05d4\u05db\u05dc", f"{SN}!$G$2:$G$10"),
    (5,  "\u05e0\u05d5\u05d1\u05dc (Nobel)",             "\u05d4\u05db\u05dc", f"{SN}!$O$2:$O$4"),
    (6,  "MWIR \u05de\u05d9\u05e0\u05d9\u05de\u05d5\u05dd",  "",  None),  # numeric, no dropdown
    (7,  "MWIR \u05de\u05e7\u05e1\u05d9\u05de\u05d5\u05dd", "",   None),
]

for row, label, default, formula1 in TOPS_FILTERS:
    lc = ws.cell(row=row, column=2, value=label)
    lc.font = Font(bold=True, size=10)
    vc = ws.cell(row=row, column=3, value=default)
    vc.fill = fill(C_LITE); vc.alignment = center()
    if formula1:
        add_dv(ws, formula1, f"C{row}")

ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 28
ws.column_dimensions["C"].width = 16

# ── Summary (rows 9-12) ───────────────────────────────────────────────
ws.cell(row=9, column=2, value="\u05e1\u05d4\"\u05db \u05e8\u05e9\u05d5\u05de\u05d5\u05ea \u05de\u05e1\u05d5\u05e0\u05e0\u05d5\u05ea").font = Font(bold=True, size=10, color=C_MID)
TECH_SH = "'\u05e2\u05e8\u05db\u05d9\u05dd \u05d8\u05db\u05e0\u05d5\u05dc\u05d5\u05d2\u05d9\u05d9\u05dd'"
ws.cell(row=9, column=3).value = f'=SUMPRODUCT(({TECH_SH}!$P$5:$P$357<>"")*1)'
ws.cell(row=9, column=3).font = Font(bold=True, size=11, color=C_DARK)

# ── Detail section headers (row 12) ──────────────────────────────────
TOPS_HDRS = [
    "\u05de\u05d6\u05d4\u05d4",         # A מזהה
    "\u05e9\u05dd \u05de\u05d5\u05e6\u05e8", # B
    "\u05e7\u05d5 \u05de\u05d5\u05e6\u05e8", # C
    "\u05d4\u05d3\u05e4\u05e1 \u05d0'", # D
    "MWIR \u05d0'",                      # E
    "LWIR \u05d0'",                      # F
    "NIR \u05d0'",                       # G
    "\u05d4\u05d3\u05e4\u05e1 \u05d1'", # H
    "MWIR \u05d1'",                      # I
    "LWIR \u05d1'",                      # J
    "NIR \u05d1'",                       # K
    "\u05e0\u05d5\u05d1\u05dc",          # L
    "\u05de\u05de\u05d5\u05e6\u05e2 MWIR", # M
    "\u05de\u05de\u05d5\u05e6\u05e2 LWIR", # N
    "\u05de\u05de\u05d5\u05e6\u05e2 \u05de\u05e9\u05d5\u05dc\u05d1", # O
]
for col_idx, h in enumerate(TOPS_HDRS, start=1):
    c = ws.cell(row=12, column=col_idx, value=h)
    c.font = hdr_font(); c.fill = fill(C_DARK); c.alignment = center(wrap=True)
ws.row_dimensions[12].height = 28

# Detail rows 13..200 from tech values via P helper
for dr in range(13, 201):
    rank = dr - 12
    for dc in range(1, 16):  # A-O (cols 1-15)
        ws.cell(row=dr, column=dc).value = (
            f'=IFERROR(INDEX({TECH_SH}!${get_column_letter(dc)}$5:${get_column_letter(dc)}$357,'
            f'SMALL({TECH_SH}!$P$5:$P$357,{rank})),"")'
        )

for c, w in {1:15,2:24,3:20,4:14,5:9,6:9,7:9,8:14,9:9,10:9,11:9,12:9,13:12,14:12,15:14}.items():
    ws.column_dimensions[get_column_letter(c)].width = w

ws.freeze_panes = "A13"
print("TechOps dashboard built.")


# ═══════════════════════════════════════════════════════════════════════
# SAVE
# ═══════════════════════════════════════════════════════════════════════
wb.save(OUT_PATH)
print(f"\nSaved: {OUT_PATH}")

# ═══════════════════════════════════════════════════════════════════════
# PHASE G — VERIFICATION
# ═══════════════════════════════════════════════════════════════════════
import zipfile, re

print("\n── Verification ─────────────────────────────────────────────────")

# 1. Can we open the zip?
with zipfile.ZipFile(OUT_PATH) as z:
    names = z.namelist()
    print(f"  ZIP entries: {len(names)}")
    # 2. Read all XML entries and check for broken tokens
    broken_tokens = re.compile(r'#REF!|#NAME\?|<f\s+t="array"', re.IGNORECASE)
    bad_files = []
    for name in names:
        if name.endswith('.xml') or name.endswith('.rels'):
            content = z.read(name).decode('utf-8', errors='replace')
            if broken_tokens.search(content):
                bad_files.append(name)
    if bad_files:
        print(f"  !! Broken tokens found in: {bad_files}")
    else:
        print("  OK: No #REF!, #NAME?, or array formula records found.")

# 3. Reload and verify structure
from openpyxl import load_workbook
wb2 = load_workbook(OUT_PATH, data_only=False)

REQUIRED_SHEETS = [
    "\u05de\u05dc\u05d0\u05d9",
    "\u05d4\u05d2\u05d3\u05e8\u05d5\u05ea",
    "\u05d4\u05e9\u05d0\u05dc\u05d5\u05ea",
    "\u05e2\u05e8\u05db\u05d9\u05dd \u05d8\u05db\u05e0\u05d5\u05dc\u05d5\u05d2\u05d9\u05d9\u05dd",
    "\u05ea\u05db\u05d5\u05dc\u05d4",
    "\u05dc\u05d5\u05d7 \u05d1\u05e7\u05e8\u05d4",
    "\u05d8\u05db\u05e0\u05d5",
]
for sn in REQUIRED_SHEETS:
    if sn in wb2.sheetnames:
        print(f"  OK: sheet '{sn}' exists.")
    else:
        print(f"  !! MISSING sheet '{sn}'")

# 4. Named ranges
REQUIRED_NR = ["PrintsTable","EnvList","LineList","ProductsTable",
                "ProductsName","ProductsSKU","FabricList","PrintsList",
                "StatusList","YesNoList","SizesList"]
for nr in REQUIRED_NR:
    if nr in wb2.defined_names:
        print(f"  OK: named range '{nr}' = {wb2.defined_names[nr].attr_text}")
    else:
        print(f"  !! MISSING named range '{nr}'")

# 5. Inventory: check row 3 formula cells
inv = wb2["\u05de\u05dc\u05d0\u05d9"]
for col, letter in [(5,"E"),(26,"Z"),(27,"AA"),(28,"AB"),(29,"AC"),(30,"AD"),(31,"AE")]:
    cell = inv.cell(row=3, column=col)
    val = str(cell.value or "")
    ok = val.startswith("=") and "#REF" not in val and "#NAME" not in val
    status = "OK" if ok else "!!"
    print(f"  {status}: מלאי {letter}3 formula: {val[:60]}...")

# 6. Settings: spot check
settings = wb2["\u05d4\u05d2\u05d3\u05e8\u05d5\u05ea"]
print(f"  OK: הגדרות A2={settings['A2'].value!r}  F2={settings['F2'].value!r}  G2={settings['G2'].value!r}")
print(f"      H2={settings['H2'].value!r}  L2={settings['L2'].value!r}  M2={settings['M2'].value!r}")

# 7. Dashboard defaults
dash = wb2["\u05dc\u05d5\u05d7 \u05d1\u05e7\u05e8\u05d4"]
print(f"  OK: Dashboard C3={dash['C3'].value!r}  C4={dash['C4'].value!r}")

print("\nDone. File is ready.")

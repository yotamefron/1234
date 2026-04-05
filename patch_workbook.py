#!/usr/bin/env python3
"""Patch ametrine_inventory.xlsx — structural fixes only."""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

PATH = "/home/user/1234/ametrine_inventory.xlsx"
wb = load_workbook(PATH)

C_DARK = "1B4332"; C_MID = "2D6A4F"; C_LITE = "D8F3DC"

def fill(h): return PatternFill("solid", fgColor=h)
def hfont(c="FFFFFF", b=True, s=10): return Font(bold=b, color=c, size=s)
def center(wrap=False): return Alignment(horizontal="center", vertical="center", wrap_text=wrap)

SN = "'\u05d4\u05d2\u05d3\u05e8\u05d5\u05ea'"  # 'הגדרות'
INV_S = "'\u05de\u05dc\u05d0\u05d9'"            # 'מלאי'
DASH_S = "'\u05dc\u05d5\u05d7 \u05d1\u05e7\u05e8\u05d4'"  # 'לוח בקרה'
TECH_S = "'\u05d8\u05db\u05e0\u05d5'"           # 'טכנו'

# ═══════════════════════════════════════════════════════════════════════
# FIX 1 — INVENTORY ROW 1 MERGED GROUPS
# ═══════════════════════════════════════════════════════════════════════
inv = wb["\u05de\u05dc\u05d0\u05d9"]

# Unmerge existing row-1 merges that need changing
for m in list(inv.merged_cells.ranges):
    if m.min_row == 1:
        inv.unmerge_cells(str(m))

def cat1(ws, c1, c2, label, fg):
    if c1 != c2:
        ws.merge_cells(start_row=1, start_column=c1, end_row=1, end_column=c2)
    cell = ws.cell(row=1, column=c1, value=label)
    cell.font = hfont(); cell.fill = fill(fg); cell.alignment = center(wrap=True)

cat1(inv,  1,  2, "\u05d6\u05d9\u05d4\u05d5\u05d9 \u05de\u05d5\u05e6\u05e8",   C_DARK)  # A:B זיהוי מוצר
cat1(inv,  3,  5, "\u05e4\u05e8\u05d8\u05d9\u05dd",                             C_MID)   # C:E פרטים
cat1(inv,  6,  7, "\u05d1\u05d3",                                                 C_MID)   # F:G בד
cat1(inv,  8,  9, "\u05d4\u05d3\u05e4\u05e1\u05d9\u05dd",                        C_MID)   # H:I הדפסים
cat1(inv, 10, 16, "\u05ea\u05e7\u05d9\u05e0\u05d5\u05ea \u05d5\u05de\u05d0\u05e4\u05d9\u05d9\u05e0\u05d9\u05dd", C_DARK)  # J:P
cat1(inv, 17, 19, "\u05d4\u05e2\u05e8\u05d5\u05ea",                              C_MID)   # Q:S הערות
cat1(inv, 20, 20, "\u05e1\u05d8\u05d8\u05d5\u05e1",                              C_DARK)  # T
cat1(inv, 21, 21, "\u05d4\u05e2\u05e8\u05d5\u05ea",                              C_DARK)  # U
cat1(inv, 22, 25, "\u05d4\u05e9\u05d0\u05dc\u05d4",                              C_MID)   # V:Y

print("FIX 1: Inventory row-1 merges corrected.")

# ═══════════════════════════════════════════════════════════════════════
# FIX 2+3 — SETTINGS PRODUCTS TABLE (exact user order, 52 products)
# ═══════════════════════════════════════════════════════════════════════
PRODUCTS = [
    ("Hobey Elite",                 "7021-00",  "OverGarment"),
    ("Gal Suit",                    "7004-01",  "OverGarment"),
    ("Sniper Assault",              "7033-00",  "OverGarment"),
    ("Poncho SRV",                  "7018-00",  "OverGarment"),
    ("Poncho Inbar",                "7030-01",  "OverGarment"),
    ("Poncho Sahar",                "7031-01",  "OverGarment"),
    ("Poncho Sahar 3D",             "",         "OverGarment"),
    ("Poncho Inbar 3D",             "",         "OverGarment"),
    ("SWAT 3-Part SUIT",            "9097-00",  "OverGarment"),
    ("May Suit",                    "7040-01",  "OverGarment"),
    ("Arctic Jacket",               "7051-00",  "OverGarment"),
    ("Arctic Overpants",            "7052-00",  "OverGarment"),
    ("Arctic Poncho SRV",           "",         "OverGarment"),
    ("Arctic Gaiters",              "",         "OverGarment"),
    ("Arctic Survival Blanket",     "",         "Blankets"),
    ("Survival Blanket",            "9021-00",  "Blankets"),
    ("Water Proof Survival Blanket","9028-00",  "Blankets"),
    ("Water Proof Sniper Blanket",  "9062-00",  "Blankets"),
    ("3D Survival Blanket",         "9029-00",  "Blankets"),
    ("IPS",                         "3007-00",  "Blankets"),
    ("Survival Blanket: Loki",      "",         "Blankets"),
    ("Operators HS",                "9001-00",  "Hide Site"),
    ("OverWatch (ATGM) HS",         "9004-01",  "Hide Site"),
    ("6X6 Hide Site",               "9095-00",  "Hide Site"),
    ("\u05e2\u05de\u05d3\u05ea \u05d7\u05d5\u05d3", "", "Hide Site"),
    ("Window Kit",                  "9083-00",  "Urban"),
    ("Balaclava",                   "8006-00",  "Uniform"),
    ("Gaiters",                     "5004-00",  "Uniform"),
    ("Headcover",                   "8030-00",  "Uniform"),
    ("Boonie hat",                  "7022-00",  "Uniform"),
    ("Vest cover",                  "7017-00",  "Uniform"),
    ("Pants",                       "",         "Uniform"),
    ("Shirt",                       "",         "Uniform"),
    ("LC79",                        "9026-00",  "Platform Hide Site"),
    ("Vehicle hide site 9*12",      "9090-00",  "Platform Hide Site"),
    ("Vehicle hide site 9*13",      "9093-00",  "Platform Hide Site"),
    ("CMM 7.5*7.5",                 "3052-00",  "Platform Hide Site"),
    ("CMM 10*10",                   "3077-05",  "Platform Hide Site"),
    ("Zodiac Shore Cover",          "3005-00",  "Platform Hide Site"),
    ("MRZR cover",                  "3003-00",  "Platform Hide Site"),
    ("Blackout hood",               "9089-00",  "Accessories"),
    ("Thermal drone landing pad",   "9077-00",  "Accessories"),
    ("Bag cover 90L",               "7009-01",  "Accessories"),
    ("Weapon Wraps",                "5001-00",  "Accessories"),
    ("Loki",                        "",         "Platform On-The-Move"),
    ("May Long Shirt",              "",         "OverGarment"),
    ("Arctic Suit",                 "",         "OverGarment"),
    ("HS Floor",                    "",         "Accessories"),
    ("\u05d7\u05dc\u05e7\u05d9\u05dd \u05e9\u05dc \u05e2\u05de\u05d3\u05d5\u05ea", "", "Accessories"),
    ("\u05e8\u05e9\u05ea \u05d4\u05e6\u05dc\u05dc\u05d4", "", "Platform Hide Site"),
    ("GMV",                         "",         "Platform Hide Site"),
    ("\u05d9\u05e8\u05d9\u05e2\u05ea \u05d4\u05d7\u05e9\u05db\u05d4", "", "Blankets"),
]  # 52 products → H2:J53

ws_set = wb["\u05d4\u05d2\u05d3\u05e8\u05d5\u05ea"]

# Clear old products area (H2:J60 to be safe)
for r in range(2, 61):
    for c in range(8, 11):
        ws_set.cell(row=r, column=c).value = None

# Write new products
for i, (name, sku, line) in enumerate(PRODUCTS, start=2):
    ws_set.cell(row=i, column=8, value=name.strip())
    ws_set.cell(row=i, column=9, value=sku.strip())
    ws_set.cell(row=i, column=10, value=line.strip())

# Also update prints list M to strip any accidental spaces
for r in range(2, 20):
    v = ws_set.cell(row=r, column=13).value
    if v:
        ws_set.cell(row=r, column=13, value=str(v).strip())

print(f"FIX 2+3: Products table rebuilt ({len(PRODUCTS)} products → H2:J{1+len(PRODUCTS)}).")

# Update named ranges for new product count (52 → rows 2-53)
def set_nr(name, ref):
    if name in wb.defined_names:
        del wb.defined_names[name]
    wb.defined_names[name] = DefinedName(name=name, attr_text=ref)

N = len(PRODUCTS) + 1  # last row = 53
set_nr("ProductsTable",  f"{SN}!$H$2:$J${N}")
set_nr("ProductsName",   f"{SN}!$H$2:$H${N}")
set_nr("ProductsSKU",    f"{SN}!$I$2:$I${N}")
print(f"  Named ranges updated: ProductsTable/Name/SKU → H2:J{N}/H2:H{N}/I2:I{N}")


# ═══════════════════════════════════════════════════════════════════════
# FIX 4 — REBUILD COMPACT DASHBOARD
# ═══════════════════════════════════════════════════════════════════════
ws_dash = wb["\u05dc\u05d5\u05d7 \u05d1\u05e7\u05e8\u05d4"]

# Wipe existing content (keep sheet, clear data/merges/validations)
for m in list(ws_dash.merged_cells.ranges):
    ws_dash.unmerge_cells(str(m))
for row in ws_dash.iter_rows():
    for cell in row:
        cell.value = None
ws_dash.data_validations.dataValidation = []

# Widths
for c, w in {1:22, 2:24, 3:5, 4:22, 5:22}.items():
    ws_dash.column_dimensions[get_column_letter(c)].width = w

# ── Row 1: Title ─────────────────────────────────────────────────────
ws_dash.merge_cells("A1:H1")
t = ws_dash["A1"]
t.value = "\u05dc\u05d5\u05d7 \u05d1\u05e7\u05e8\u05d4 \u2013 \u05de\u05dc\u05d0\u05d9 \u05de\u05d3\u05d2\u05d9\u05de\u05d9 \u05e9\u05d9\u05d5\u05d5\u05e7"
t.font = Font(bold=True, size=15, color=C_DARK); t.alignment = center()
ws_dash.row_dimensions[1].height = 26

# ── Rows 4-9: Filters (two blocks) ───────────────────────────────────
ws_dash.row_dimensions[3].height = 8  # spacer

LEFT = [
    (4, "\u05e1\u05d1\u05d9\u05d1\u05d4:",                "\u05d4\u05db\u05dc"),  # סביבה
    (5, "\u05e7\u05d5 \u05de\u05d5\u05e6\u05e8:",         "\u05d4\u05db\u05dc"),  # קו מוצר
    (6, "\u05de\u05d5\u05e6\u05e8:",                       ""),                    # מוצר (free text)
    (7, "\u05e1\u05d5\u05d2 \u05d1\u05d3:",               "\u05d4\u05db\u05dc"),  # סוג בד
    (8, "\u05e1\u05d8\u05d8\u05d5\u05e1:",                "\u05d4\u05db\u05dc"),  # סטטוס
    (9, "\u05d2\u05e8\u05e1\u05d4 \u05de\u05d9\u05d5\u05d7\u05d3\u05ea(\u05e4\u05d9\u05ea\u05d5\u05d7):", "\u05d4\u05db\u05dc"),  # גרסה מיוחדת
]
RIGHT = [
    (4, "\u05ea\u05e7\u05d9\u05e0\u05d5\u05ea \u05d5\u05d9\u05d6\u05d5\u05d0\u05dc\u05d9\u05ea:", "\u05d4\u05db\u05dc"),  # תקינות ויזואלית
    (5, "\u05ea\u05e7\u05d9\u05e0\u05d5\u05ea \u05ea\u05e8\u05de\u05d9\u05ea:",                   "\u05d4\u05db\u05dc"),  # תקינות תרמית
    (6, "\u05d2\u05e8\u05e1\u05d4 \u05e1\u05d5\u05e4\u05d9\u05ea:",                              "\u05d4\u05db\u05dc"),  # גרסה סופית
    (7, "\u05de\u05d9\u05d3\u05d4:",                                                              "\u05d4\u05db\u05dc"),  # מידה
    (8, "\u05de\u05d3\u05d9\u05e0\u05ea \u05d4\u05e9\u05d0\u05dc\u05d4:",                        ""),                   # מדינת השאלה (text)
    (9, "\u05d2\u05e8\u05e1\u05ea \u05de\u05d5\u05e6\u05e8:",                                    ""),                   # גרסת מוצר (text)
]

def add_dv2(ws, f1, sqref):
    dv = DataValidation(type="list", formula1=f1, allow_blank=True, showDropDown=False)
    dv.sqref = sqref
    ws.add_data_validation(dv)

# Left block labels (col A) and values (col B)
LEFT_VALIDATIONS = {
    4: f"{SN}!$F$2:$F$8",   # env
    5: f"{SN}!$G$2:$G$10",  # line
    # 6: free text
    7: f"{SN}!$L$2:$L$18",  # fabric
    8: f"{SN}!$N$2:$N$6",   # status
    9: f"{SN}!$O$2:$O$4",   # yes/no
}
RIGHT_VALIDATIONS = {
    4: f"{SN}!$O$2:$O$4",  # visual
    5: f"{SN}!$O$2:$O$4",  # thermal
    6: f"{SN}!$O$2:$O$4",  # final
    7: f"{SN}!$P$2:$P$9",  # size
    # 8,9: free text
}

for row, label, default in LEFT:
    lc = ws_dash.cell(row=row, column=1, value=label)
    lc.font = Font(bold=True, size=10, color=C_DARK)
    lc.alignment = Alignment(horizontal="right", vertical="center")
    vc = ws_dash.cell(row=row, column=2, value=default)
    vc.fill = fill(C_LITE); vc.alignment = center(); vc.font = Font(size=10)
    if row in LEFT_VALIDATIONS:
        add_dv2(ws_dash, LEFT_VALIDATIONS[row], f"B{row}")

for row, label, default in RIGHT:
    lc = ws_dash.cell(row=row, column=4, value=label)
    lc.font = Font(bold=True, size=10, color=C_DARK)
    lc.alignment = Alignment(horizontal="right", vertical="center")
    vc = ws_dash.cell(row=row, column=5, value=default)
    vc.fill = fill(C_LITE); vc.alignment = center(); vc.font = Font(size=10)
    if row in RIGHT_VALIDATIONS:
        add_dv2(ws_dash, RIGHT_VALIDATIONS[row], f"E{row}")

# ── Rows 10-14: Summary ───────────────────────────────────────────────
ws_dash.row_dimensions[10].height = 6  # small spacer row

INV = "'מלאי'"
SUMMARY_L = [
    (11, "\u05e1\u05d4\"\u05db \u05de\u05d5\u05e6\u05e8\u05d9\u05dd \u05de\u05e1\u05d5\u05e0\u05e0\u05d9\u05dd",
         f"=SUMPRODUCT(({INV}!$AD$3:$AD$357<>\"\")*1)"),
    (12, "\u05d1\u05de\u05dc\u05d0\u05d9",
         f"=SUMPRODUCT(({INV}!$AD$3:$AD$357<>\"\")*({INV}!$T$3:$T$357=\"\u05d1\u05de\u05dc\u05d0\u05d9\"))"),
    (13, "\u05de\u05d5\u05e9\u05d0\u05dc",
         f"=SUMPRODUCT(({INV}!$AD$3:$AD$357<>\"\")*({INV}!$T$3:$T$357=\"\u05de\u05d5\u05e9\u05d0\u05dc\"))"),
    (14, "\u05d1\u05d7\u05d3\u05e8 \u05ea\u05e6\u05d5\u05d2\u05d4",
         f"=SUMPRODUCT(({INV}!$AD$3:$AD$357<>\"\")*({INV}!$T$3:$T$357=\"\u05d1\u05d7\u05d3\u05e8 \u05ea\u05e6\u05d5\u05d2\u05d4\"))"),
]
SUMMARY_R = [
    (11, "\u05d1\u05ea\u05d9\u05e7 \u05d4\u05d3\u05d2\u05de\u05d4",
         f"=SUMPRODUCT(({INV}!$AD$3:$AD$357<>\"\")*({INV}!$T$3:$T$357=\"\u05d1\u05ea\u05d9\u05e7 \u05d4\u05d3\u05d2\u05de\u05d4\"))"),
    (12, "\u05ea\u05e7\u05d9\u05df \u05d5\u05d9\u05d6\u05d5\u05d0\u05dc\u05d9 + \u05ea\u05e8\u05de\u05d9",
         f"=SUMPRODUCT(({INV}!$AD$3:$AD$357<>\"\")*({INV}!$J$3:$J$357=\"\u05db\u05df\")*({INV}!$L$3:$L$357=\"\u05db\u05df\"))"),
    (13, "\u05d2\u05e8\u05e1\u05d4 \u05de\u05d9\u05d5\u05d7\u05d3\u05ea",
         f"=SUMPRODUCT(({INV}!$AD$3:$AD$357<>\"\")*({INV}!$Q$3:$Q$357=\"\u05db\u05df\"))"),
    (14, "\u05d2\u05e8\u05e1\u05d4 \u05e1\u05d5\u05e4\u05d9\u05ea",
         f"=SUMPRODUCT(({INV}!$AD$3:$AD$357<>\"\")*({INV}!$O$3:$O$357=\"\u05db\u05df\"))"),
]

for row, label, formula in SUMMARY_L:
    ws_dash.cell(row=row, column=1, value=label).font = Font(size=10, bold=True)
    c = ws_dash.cell(row=row, column=2, value=formula)
    c.font = Font(bold=True, size=12, color=C_DARK); c.alignment = center()

for row, label, formula in SUMMARY_R:
    ws_dash.cell(row=row, column=4, value=label).font = Font(size=10, bold=True)
    c = ws_dash.cell(row=row, column=5, value=formula)
    c.font = Font(bold=True, size=12, color=C_DARK); c.alignment = center()

# ── Row 15: spacer; Row 16: section title ────────────────────────────
ws_dash.row_dimensions[15].height = 6
ws_dash.merge_cells("A16:H16")
t16 = ws_dash["A16"]
t16.value = "\u05e4\u05d9\u05e8\u05d5\u05d8 \u05de\u05d5\u05e6\u05e8\u05d9\u05dd \u05de\u05e1\u05d5\u05e0\u05e0\u05d9\u05dd"  # פירוט מוצרים מסוננים
t16.font = hfont(s=11); t16.fill = fill(C_DARK); t16.alignment = center()

# ── Row 17: Detail headers ────────────────────────────────────────────
DETAIL_HDR = [
    "\u05e7\u05d5 \u05de\u05d5\u05e6\u05e8",   # קו מוצר
    "\u05e9\u05dd \u05de\u05d5\u05e6\u05e8",    # שם מוצר
    "\u05d2\u05e8\u05e1\u05d4",                  # גרסה
    "\u05de\u05d6\u05d4\u05d4",                  # מזהה
    "\u05d4\u05d3\u05e4\u05e1 \u05d0'",          # הדפס א'
    "\u05d4\u05d3\u05e4\u05e1 \u05d1'",          # הדפס ב'
    "\u05e1\u05d8\u05d8\u05d5\u05e1",            # סטטוס
    "\u05d4\u05e2\u05e8\u05d5\u05ea",            # הערות
]
for ci, h in enumerate(DETAIL_HDR, start=1):
    c = ws_dash.cell(row=17, column=ci, value=h)
    c.font = hfont(s=9); c.fill = fill(C_MID); c.alignment = center(wrap=True)
ws_dash.row_dimensions[17].height = 22

# ── Rows 18-217: Detail data via SMALL+INDEX ──────────────────────────
# Inventory columns: A=1,B=2,C=3,D=4,H=8,I=9,T=20,U=21
INV_COLS = {1:"A",2:"B",3:"C",4:"D",5:"H",6:"I",7:"T",8:"U"}
for dr in range(18, 218):
    rank = dr - 17  # 1,2,...
    for dc, ic in INV_COLS.items():
        ws_dash.cell(row=dr, column=dc).value = (
            f'=IFERROR(INDEX({INV}!${ic}$3:${ic}$357,'
            f'SMALL({INV}!$AD$3:$AD$357,{rank})),"")'
        )

# Detail column widths
for c, w in {1:16,2:22,3:12,4:14,5:14,6:14,7:14,8:24}.items():
    ws_dash.column_dimensions[get_column_letter(c)].width = w

ws_dash.freeze_panes = "A18"
print("FIX 4: Dashboard rebuilt (compact, rows 4-9 filters, 11-14 summary, 17+ detail).")


# ═══════════════════════════════════════════════════════════════════════
# FIX 5 — REWRITE AD FORMULAS IN INVENTORY (new dashboard filter refs)
# ═══════════════════════════════════════════════════════════════════════
# New filter cell layout in dashboard:
#   B4=env, B5=line, B6=product(text), B7=fabric, B8=status, B9=dev
#   E4=visual, E5=thermal, E6=final, E7=size, E8=country(text), E9=version(text)
MAX_ROW = 357
inv = wb["\u05de\u05dc\u05d0\u05d9"]

for r in range(3, MAX_ROW + 1):
    # AD: dashboard filter helper
    inv.cell(row=r, column=30).value = (
        f'=IF(AND('
        f'D{r}<>"",'
        f'OR({DASH_S}!$B$4="\u05d4\u05db\u05dc",{DASH_S}!$B$4="",'
        f'Z{r}={DASH_S}!$B$4,AA{r}={DASH_S}!$B$4,AB{r}={DASH_S}!$B$4,AC{r}={DASH_S}!$B$4),'
        f'OR({DASH_S}!$B$5="\u05d4\u05db\u05dc",{DASH_S}!$B$5="",A{r}={DASH_S}!$B$5),'
        f'OR({DASH_S}!$B$6="",ISNUMBER(SEARCH({DASH_S}!$B$6,B{r}))),'
        f'OR({DASH_S}!$B$7="\u05d4\u05db\u05dc",{DASH_S}!$B$7="",'
        f'F{r}={DASH_S}!$B$7,G{r}={DASH_S}!$B$7),'
        f'OR({DASH_S}!$B$8="\u05d4\u05db\u05dc",{DASH_S}!$B$8="",T{r}={DASH_S}!$B$8),'
        f'OR({DASH_S}!$B$9="\u05d4\u05db\u05dc",{DASH_S}!$B$9="",Q{r}={DASH_S}!$B$9),'
        f'OR({DASH_S}!$E$4="\u05d4\u05db\u05dc",{DASH_S}!$E$4="",J{r}={DASH_S}!$E$4),'
        f'OR({DASH_S}!$E$5="\u05d4\u05db\u05dc",{DASH_S}!$E$5="",L{r}={DASH_S}!$E$5),'
        f'OR({DASH_S}!$E$6="\u05d4\u05db\u05dc",{DASH_S}!$E$6="",O{r}={DASH_S}!$E$6),'
        f'OR({DASH_S}!$E$7="\u05d4\u05db\u05dc",{DASH_S}!$E$7="",R{r}={DASH_S}!$E$7),'
        f'OR({DASH_S}!$E$8="",ISNUMBER(SEARCH({DASH_S}!$E$8,W{r}))),'
        f'OR({DASH_S}!$E$9="",ISNUMBER(SEARCH({DASH_S}!$E$9,C{r})))'
        f'),ROW()-2,"")'
    )
    # AE unchanged — keep existing techops ref

print(f"FIX 5: AD formulas rewritten for rows 3-{MAX_ROW} (new dashboard filter refs).")

# ═══════════════════════════════════════════════════════════════════════
# FIX 6 — LOANS: hide helper col A (row counter), keep sheet readable
# ═══════════════════════════════════════════════════════════════════════
ws_loan = wb["\u05d4\u05e9\u05d0\u05dc\u05d5\u05ea"]
# The AGGREGATE formula in col B is the real driver — col A is just a seq #.
# Col A's formula is already minimal (=IF(B5="","",N)). Keep as-is, no hide needed.
# Just verify col B formula is intact:
b5 = ws_loan.cell(row=5, column=2).value
ok = b5 and "AGGREGATE" in str(b5)
print(f"FIX 6: Loans B5 AGGREGATE formula {'OK' if ok else 'MISSING'}: {str(b5)[:60] if b5 else 'None'}")

# ═══════════════════════════════════════════════════════════════════════
# SAVE + VERIFY
# ═══════════════════════════════════════════════════════════════════════
wb.save(PATH)
print(f"\nSaved: {PATH}")

# Reload + verify
import zipfile, re
from openpyxl import load_workbook as lw2

wb2 = lw2(PATH)

# XML scan
with zipfile.ZipFile(PATH) as z:
    bad = [n for n in z.namelist()
           if n.endswith('.xml') and
           re.search(r'#REF!|#NAME\?|<f\s+t="array"', z.read(n).decode('utf-8','replace'))]
print("XML scan:", "CLEAN" if not bad else f"ISSUES in {bad}")

# Inventory row-1 groups
inv2 = wb2["\u05de\u05dc\u05d0\u05d9"]
print("Inventory row-1 merges:", sorted(str(m) for m in inv2.merged_cells.ranges if 'A1' in str(m) or '1:' in str(m)))

# Products table spot-check
ws2 = wb2["\u05d4\u05d2\u05d3\u05e8\u05d5\u05ea"]
print(f"Settings H2={ws2['H2'].value!r}  H11={ws2.cell(11,8).value!r}  H53={ws2.cell(53,8).value!r}")

# Named ranges
for nr in ["ProductsName","ProductsSKU","ProductsTable"]:
    print(f"  {nr}: {wb2.defined_names[nr].attr_text if nr in wb2.defined_names else 'MISSING'}")

# Dashboard filter cells
dash2 = wb2["\u05dc\u05d5\u05d7 \u05d1\u05e7\u05e8\u05d4"]
print(f"Dashboard B4={dash2['B4'].value!r} B5={dash2['B5'].value!r} E4={dash2['E4'].value!r}")
print(f"Dashboard summary B11={dash2['B11'].value[:50] if dash2['B11'].value else None}")

# AD formula sample
inv3 = wb2["\u05de\u05dc\u05d0\u05d9"]
ad3 = inv3.cell(3, 30).value
print(f"AD3 formula (first 90 chars): {str(ad3)[:90]}")

print("\nDone.")

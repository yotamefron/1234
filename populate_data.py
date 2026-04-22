"""
Populate the new ametrine_inventory.xlsx with data from the old marketing_inventory file.

Steps:
1. Read all data from the old Excel file.
2. Run build_excel.py to generate a fresh ametrine_inventory.xlsx.
3. Open the generated file and populate it with real data.
4. Save the file.
"""

import subprocess
import sys
import openpyxl

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
OLD_FILE = "marketing_inventory 2026 לשליחה.xlsx"
NEW_FILE = "ametrine_inventory.xlsx"
PASSWORD = "1998"

# Old inventory (מלאי) column mapping: old_col (1-based) -> new_col (1-based)
# Old data starts at row 3, new data starts at row 2.
# Column 5 (מק"ט) is skipped — formula in new file.
INV_COL_MAP = {
    1: 1,    # קו מוצר → A
    2: 2,    # שם מוצר → B
    3: 3,    # גרסה → C
    4: 4,    # מזהה ייחודי → D
    # 5: SKIP (formula in col E)
    6: 7,    # סוג בד 1 → G
    7: 8,    # סוג בד 2 → H
    8: 13,   # הדפס צד א' → M
    9: 14,   # הדפס צד ב' → N
    10: 15,  # תקין ויזואלית א' → O
    11: 16,  # תקין ויזואלית ב' → P
    12: 21,  # תקין תרמית A' → U
    13: 22,  # תקין תרמית ב' → V
    14: 23,  # פיצ'רים תקינים → W
    15: 11,  # גרסה סופית → K
    16: 12,  # הערות גרסה/פיצ'רים → L
    17: 10,  # גרסה מיוחדת(פיתוח) → J
    18: 9,   # מידה → I
    19: 30,  # פער בתכולה → AD
    20: 24,  # סטטוס → X
    21: 29,  # הערות → AC
    22: 25,  # מושאל ל → Y
    23: 26,  # מדינה → Z
    24: 27,  # אחראי מהחברה → AA
    25: 28,  # החזרה משוערת → AB
}

# New inventory formula columns that must NOT be overwritten (1-based)
INV_FORMULA_COLS = {5, 17, 18, 19, 20, 31, 32, 33, 34, 35, 36, 37}  # E, Q-T, AE-AK

# Old desired inventory (מלאי רצוי) column mapping: old_col -> new_col
DI_COL_MAP = {
    1: 1,    # שם מוצר → A
    2: 6,    # רצוי חדר תצוגה → F
    3: 7,    # רצוי בתיק הדגמות → G
    4: 8,    # רצוי השאלות → H
    5: 3,    # סביבה צד אחד → C
    6: 4,    # סביבה צד שני → D
    7: 5,    # סוג בד → E
    8: 15,   # הערות → O
}

# New desired inventory formula columns that must NOT be overwritten (1-based)
DI_FORMULA_COLS = {9, 10, 11, 12, 13, 14}  # I-N


def clean_value(val):
    """Return None for values that should be skipped (0, '#REF!', empty)."""
    if val is None:
        return None
    if val == 0:
        return None
    if isinstance(val, str):
        s = val.strip()
        if s == "" or s == "#REF!":
            return None
        return s
    return val


def main():
    # -----------------------------------------------------------------
    # Step 1: Read all data from the old file
    # -----------------------------------------------------------------
    print(f"Reading old file: {OLD_FILE}")
    old_wb = openpyxl.load_workbook(OLD_FILE, data_only=True)

    # --- Read old inventory data ---
    old_inv = old_wb["מלאי"]
    inv_data = []
    for r in range(3, old_inv.max_row + 1):
        # Check if columns 1-4 are all empty
        key_vals = [old_inv.cell(row=r, column=c).value for c in range(1, 5)]
        if all(v is None or (isinstance(v, str) and v.strip() == "") for v in key_vals):
            continue  # Skip empty rows

        row_data = {}
        for old_col, new_col in INV_COL_MAP.items():
            val = clean_value(old_inv.cell(row=r, column=old_col).value)
            if val is not None:
                row_data[new_col] = val
        inv_data.append(row_data)

    print(f"  Inventory rows read: {len(inv_data)}")

    # --- Read old desired inventory data ---
    old_di = old_wb["מלאי רצוי"]
    di_data = []
    for r in range(2, old_di.max_row + 1):
        # Skip rows where product name (col 1) is empty
        product = old_di.cell(row=r, column=1).value
        if product is None or (isinstance(product, str) and product.strip() == ""):
            continue

        row_data = {}
        for old_col, new_col in DI_COL_MAP.items():
            val = clean_value(old_di.cell(row=r, column=old_col).value)
            if val is not None:
                row_data[new_col] = val
        di_data.append(row_data)

    print(f"  Desired inventory rows read: {len(di_data)}")
    old_wb.close()

    # -----------------------------------------------------------------
    # Step 2: Run build_excel.py to generate a fresh ametrine_inventory.xlsx
    # -----------------------------------------------------------------
    print("\nRunning build_excel.py...")
    result = subprocess.run(
        [sys.executable, "build_excel.py"],
        capture_output=True,
        text=True,
    )
    if result.returncode != 0:
        print(f"ERROR: build_excel.py failed:\n{result.stderr}")
        sys.exit(1)
    print("  build_excel.py completed successfully.")

    # -----------------------------------------------------------------
    # Step 3: Open generated file and populate with data
    # -----------------------------------------------------------------
    print(f"\nOpening {NEW_FILE} for writing...")
    wb = openpyxl.load_workbook(NEW_FILE)

    # --- Populate inventory sheet ---
    ws_inv = wb["מלאי"]

    # Unprotect the sheet
    ws_inv.protection.sheet = False

    print(f"  Writing {len(inv_data)} inventory rows...")
    for i, row_data in enumerate(inv_data):
        new_row = i + 2  # Data starts at row 2 in new file
        for new_col, val in row_data.items():
            if new_col in INV_FORMULA_COLS:
                continue  # Don't overwrite formula cells
            ws_inv.cell(row=new_row, column=new_col, value=val)

    # Re-protect the sheet
    ws_inv.protection.sheet = True
    ws_inv.protection.password = PASSWORD
    ws_inv.protection.enable()

    # --- Populate desired inventory sheet ---
    ws_di = wb["מלאי רצוי"]

    # Unprotect the sheet
    ws_di.protection.sheet = False

    print(f"  Writing {len(di_data)} desired inventory rows...")
    for i, row_data in enumerate(di_data):
        new_row = i + 2  # Data starts at row 2 in new file
        for new_col, val in row_data.items():
            if new_col in DI_FORMULA_COLS:
                continue  # Don't overwrite formula cells
            ws_di.cell(row=new_row, column=new_col, value=val)

    # Re-protect the sheet
    ws_di.protection.sheet = True
    ws_di.protection.password = PASSWORD
    ws_di.protection.enable()

    # -----------------------------------------------------------------
    # Step 4: Save the file
    # -----------------------------------------------------------------
    print(f"\nSaving {NEW_FILE}...")
    wb.save(NEW_FILE)
    wb.close()
    print("Done!")

    # -----------------------------------------------------------------
    # Step 5: Verify the data was written correctly
    # -----------------------------------------------------------------
    print("\n--- Verification ---")
    wb_check = openpyxl.load_workbook(NEW_FILE, data_only=False)

    ws_check = wb_check["מלאי"]
    inv_count = 0
    for r in range(2, 500):
        if ws_check.cell(row=r, column=4).value is not None:  # Check unique ID col D
            inv_count += 1
    print(f"  Inventory rows with data (col D): {inv_count}")

    # Spot check first row
    print(f"  Row 2: A={ws_check.cell(row=2, column=1).value}, "
          f"B={ws_check.cell(row=2, column=2).value}, "
          f"D={ws_check.cell(row=2, column=4).value}, "
          f"E(formula)={ws_check.cell(row=2, column=5).value}, "
          f"G={ws_check.cell(row=2, column=7).value}, "
          f"M={ws_check.cell(row=2, column=13).value}")
    # Check a formula column is still a formula
    e2 = ws_check.cell(row=2, column=5).value
    assert isinstance(e2, str) and e2.startswith("="), f"E2 should be a formula, got: {repr(e2)}"
    print(f"  E2 formula intact: YES")

    # Check desired inventory
    ws_di_check = wb_check["מלאי רצוי"]
    di_count = 0
    for r in range(2, 200):
        if ws_di_check.cell(row=r, column=1).value is not None:
            di_count += 1
    print(f"  Desired inventory rows with data: {di_count}")

    # Spot check first desired row
    print(f"  DI Row 2: A={ws_di_check.cell(row=2, column=1).value}, "
          f"C={ws_di_check.cell(row=2, column=3).value}, "
          f"F={ws_di_check.cell(row=2, column=6).value}, "
          f"I(formula)={ws_di_check.cell(row=2, column=9).value}")
    # Check formula column I is still a formula
    i2 = ws_di_check.cell(row=2, column=9).value
    assert isinstance(i2, str) and i2.startswith("="), f"I2 should be a formula, got: {repr(i2)}"
    print(f"  I2 formula intact: YES")

    # Check protection
    assert ws_check.protection.sheet, "Inventory sheet should be protected"
    assert ws_di_check.protection.sheet, "Desired inventory sheet should be protected"
    print(f"  Sheet protection: ACTIVE on both sheets")

    wb_check.close()
    print("\nAll verifications passed!")


if __name__ == "__main__":
    main()
